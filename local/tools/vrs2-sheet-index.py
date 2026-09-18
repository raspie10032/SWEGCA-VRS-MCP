#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""스프레드시트를 스키마 단위 경험으로 색인한다 (2026-09-18, 개선 ⑥).

파일이 아니라 **스키마**(프로젝트 × 시트 헤더 서명)가 기록 하나다: 헤더·열 수·열 형(숫자/날짜/문자 — 값은 절대
안 넣는다, 회사 데이터)·같은 구조의 파일 목록(경로·시트·행 수·수정일)·폴더. 월별 output 처럼 같은 구조가
반복되면 파일 700개가 기록 수십 개가 된다. cue 는 헤더 낱말·파일명·시트명·폴더명(본문에서 데몬이 뽑는다).
outcome 은 pending — 길 찾기 기억이지 검증된 경험이 아니다(판정이 evidence 로 가리킬 때 경험이 된다).
revision = 파일 목록+수정일+행 수 해시 → 달이 바뀌어 파일이 늘면 supersedes 로 갱신(매니페스트 `vrs2-sheet-manifest.json`).

    vrs2-venv python vrs2-sheet-index.py <폴더> [--label 이름] [--project 슬러그] [--dry-run] [--max-files 12]

read_only 시트는 iter_rows 전에 reset_dimensions() — 파일의 <dimension> 이 A1 로 잘못 박힌 판이 있다
(SQLITE 메모리 feedback_openpyxl_readonly_dimension).
"""
import argparse
import csv
import datetime as dt
import hashlib
import importlib.util
import io
import json
import os
import re
import sys
import time

SKIP_DIRS = {".git", "__pycache__", "node_modules", "archive", "backups", "backup", ".venv", "venv"}
SKIP_PREFIX = ("_deleted_backup", "~$")
EXTS = (".xlsx", ".xlsm", ".csv")
MANIFEST = r"C:/Users/asm/mcp/vrs2-sheet-manifest.json"
PRODUCE = r"C:/Users/asm/mcp/vrs2-produce.py"      # 데몬 접속 헬퍼만 빌린다(ensure_daemon + PY)
SAMPLE_ROWS = 30
MAX_ROWS = 2000            # 행 수는 힌트다: 이 너머는 세지 않고 '2000+' (큰 파일에서 read_only 순회가 분 단위)


def kind_of(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return "논리"
    if isinstance(value, (int, float)):
        return "숫자"
    if isinstance(value, (dt.date, dt.datetime)):
        return "날짜"
    s = str(value).strip()
    if re.fullmatch(r"-?\d{1,3}(,\d{3})*(\.\d+)?|-?\d+(\.\d+)?", s):
        return "숫자"
    if re.fullmatch(r"\d{4}[-./]\d{1,2}[-./]\d{1,2}.*", s):
        return "날짜"
    return "문자"


def sheets_of_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            ws.reset_dimensions()                       # 잘못 박힌 <dimension> 함정
            rows = ws.iter_rows(values_only=True)
            header = None; count = 0; kinds = {}
            for row in rows:
                if header is None:
                    if row is None or all(v is None or str(v).strip() == "" for v in row):
                        continue
                    header = [str(v).strip() if v is not None else "" for v in row]
                    while header and header[-1] == "":
                        header.pop()
                    continue
                count += 1
                if count <= SAMPLE_ROWS:
                    for k, v in zip(header, row):
                        t = kind_of(v)
                        if t:
                            kinds.setdefault(k, {}).setdefault(t, 0)
                            kinds[k][t] += 1
                if count >= MAX_ROWS:
                    count = f"{MAX_ROWS}+"
                    break
            if header:
                yield ws.title, header, count, kinds
    finally:
        wb.close()


def sheets_of_csv(path):
    raw = io.open(path, "rb").read(200000)
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp949"):
        try:
            text = raw.decode(enc); break
        except UnicodeDecodeError:
            continue
    if text is None:
        return
    reader = csv.reader(io.StringIO(text))
    header = None; count = 0; kinds = {}
    for row in reader:
        if header is None:
            if not row or all(not c.strip() for c in row):
                continue
            header = [c.strip() for c in row]
            while header and header[-1] == "":
                header.pop()
            continue
        count += 1
        if count <= SAMPLE_ROWS:
            for k, v in zip(header, row):
                t = kind_of(v)
                if t:
                    kinds.setdefault(k, {}).setdefault(t, 0); kinds[k][t] += 1
    if header:
        # 200 KB 만 읽었으니 행 수는 하한이다
        yield "csv", header, count, kinds


def walk(root):
    for base, dirs, files in os.walk(root):
        dirs[:] = [d for d in dirs if d not in SKIP_DIRS and not d.startswith(SKIP_PREFIX)]
        for name in files:
            if name.lower().endswith(EXTS) and not name.startswith(SKIP_PREFIX):
                yield os.path.join(base, name)


def header_like(header):
    """헤더답지 않은 첫 줄(빈 칸·숫자 일련번호가 대부분인 원본 시트)은 스키마로 안 센다."""
    names = [h for h in header if h and not re.fullmatch(r"[\d.\s-]+", h)]
    return len(names) >= 2 and len(names) >= len(header) * 0.5


def signature(header):
    norm = [re.sub(r"\s+", "", h.casefold()) for h in header]
    return hashlib.sha1("|".join(norm).encode("utf-8")).hexdigest()[:8]


def collect(root):
    schemas = {}
    errors = []
    for path in walk(root):
        rel = os.path.relpath(path, root).replace("\\", "/")
        try:
            mtime = time.strftime("%Y-%m-%d", time.localtime(os.path.getmtime(path)))
            gen = sheets_of_csv(path) if path.lower().endswith(".csv") else sheets_of_xlsx(path)
            for title, header, count, kinds in gen:
                if not header_like(header):
                    continue
                sig = signature(header)
                s = schemas.setdefault(sig, dict(header=header, files=[], kinds={}))
                s["files"].append(dict(file=rel, sheet=title, rows=count, mtime=mtime))
                for k, ks in kinds.items():
                    tgt = s["kinds"].setdefault(k, {})
                    for t, n in ks.items():
                        tgt[t] = tgt.get(t, 0) + n
        except Exception as failure:
            errors.append((rel, repr(failure)[:120]))
    return schemas, errors


def render(project, sig, schema, max_files):
    header = schema["header"]
    files = sorted(schema["files"], key=lambda f: (f["mtime"], f["file"]), reverse=True)
    folders = sorted({f["file"].rsplit("/", 1)[0] if "/" in f["file"] else "." for f in files})
    def kind_label(col):
        ks = schema["kinds"].get(col) or {}
        return max(ks, key=ks.get) if ks else "빈"
    lines = [f"시트 스키마 {project} · {sig}",
             f"헤더({len(header)}열): " + " | ".join(h or "(빈)" for h in header),
             "열 형: " + ", ".join(f"{h or '(빈)'}={kind_label(h)}" for h in header[:40]),
             f"파일 {len(files)}개 (최근순, 시트·행 수·수정일): " + "; ".join(
                 f"{f['file']}!{f['sheet']} ({f['rows']}행, {f['mtime']})" for f in files[:max_files])
             + (f" … 외 {len(files) - max_files}개" if len(files) > max_files else ""),
             "폴더: " + ", ".join(folders[:12]),
             "값은 색인하지 않는다(헤더·형·파일 목록만). 찾을 때 묻는 말: 이 컬럼 있는 파일이 뭐냐 · 그 달 파일 어디 있나 · 헤더가 어떻게 생겼나"]
    text = "\n".join(lines)
    revision = hashlib.sha1(json.dumps([[f["file"], f["sheet"], f["rows"], f["mtime"]] for f in files], ensure_ascii=False).encode()).hexdigest()[:12]
    return text, revision, files, folders


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("root")
    ap.add_argument("--label", default=None, help="출처 이름(기본: 폴더 이름)")
    ap.add_argument("--project", default=None, help="기억 프로젝트 슬러그(기본: 훅과 같은 규칙 — 그 폴더의 memory/ 를 가진 조상 프로젝트)")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--max-files", type=int, default=12)
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    sys.path.insert(0, os.path.join(os.path.expanduser("~"), ".claude", "hooks"))
    from project_dir import resolve
    label = a.label or os.path.basename(os.path.normpath(a.root))
    if a.project is None:
        _slug, memory = resolve(a.root)
        a.project = os.path.basename(os.path.dirname(memory))     # 정션이면 실제 프로젝트(SETTLEMENT → SQLITE)
    print("프로젝트", a.project, "| 출처 라벨", label)
    schemas, errors = collect(a.root)
    print(f"{a.root}: 스키마 {len(schemas)}개, 파일 {sum(len(s['files']) for s in schemas.values())}개, 오류 {len(errors)}개")
    for rel, err in errors[:5]:
        print("  오류", rel, err)
    try:
        manifest = json.load(io.open(MANIFEST, encoding="utf-8")) if os.path.exists(MANIFEST) else {}
    except ValueError:
        manifest = {}
    client = None
    if not a.dry_run:
        spec = importlib.util.spec_from_file_location("vrs2_produce", PRODUCE); mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
        sys.path.insert(0, mod.SRC)
        from swegca_vrs2.loopback import ensure_daemon
        client = ensure_daemon(mod.STATE, allow_ingest=True, python=mod.PY)
    added = updated = skipped = 0
    try:
        for sig, schema in sorted(schemas.items(), key=lambda kv: -len(kv[1]["files"])):
            text, revision, files, folders = render(label, sig, schema, a.max_files)
            revision = hashlib.sha1((revision + "|" + a.project).encode()).hexdigest()[:12]   # 프로젝트가 바뀌어도 새 판
            source = f"sheet:{label}/{sig}"
            prev = manifest.get(source)
            if a.dry_run:
                print("\n" + text[:700] + ("…" if len(text) > 700 else ""))
                continue
            if prev and prev["revision"] == revision:
                skipped += 1
                continue
            args = dict(request_id=f"sheet:{source}@{revision}", text=text[:60000], source=source, revision=revision,
                        outcome="pending", cues=[],
                        metadata=dict(kind="sheet_schema", project=a.project, files=len(files), columns=len(schema["header"]),
                                      headers=schema["header"][:60], folders=folders[:12]))
            if prev:
                args["supersedes"] = prev["episode"]
            try:
                out = client.request("ingest", **args)
            except Exception as failure:
                if "supersedes" in args:
                    args.pop("supersedes"); out = client.request("ingest", **args)
                else:
                    print("  실패", source, repr(failure)[:120]); continue
            manifest[source] = dict(revision=revision, episode=out.get("episode_id"))
            added += 0 if prev else 1; updated += 1 if prev else 0
    finally:
        if client is not None:
            client.close()
    if not a.dry_run:
        io.open(MANIFEST, "w", encoding="utf-8").write(json.dumps(manifest, ensure_ascii=False, indent=0))
        print(f"넣음 {added} 갱신 {updated} 건너뜀 {skipped}")


if __name__ == "__main__":
    main()
